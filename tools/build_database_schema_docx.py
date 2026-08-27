import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "REVAMPED_DATABASE_SCHEMA.md"
OUTPUT = ROOT / "docs" / "School_Management_System_Database_Schema.docx"

BLUE = "2E74B5"
NAVY = "1F4D78"
PALE = "E8EEF5"
LIGHT = "F5F8FB"
GRID = "B8C6D1"
TEXT = "25313C"
MUTED = "667788"
WHITE = "FFFFFF"

TABLE_PURPOSES = {
    "USERS": "Authentication and account-level authorization.",
    "STAFF_PROFILES": "Normalized identity and employment profile for administrators, registrars, and guards.",
    "POSITIONS": "Controlled staff position catalog.",
    "DEPARTMENTS": "Academic and organizational department catalog.",
    "SCHOOL_YEARS": "School-year calendar and enrollment window configuration.",
    "GRADE_LEVELS": "Education-level catalog; distinct from student marks.",
    "SECTIONS": "A class grouping for a grade level within one school year.",
    "TEACHERS": "Permanent teacher identity and employment profile.",
    "SECTION_TEACHERS": "Teacher-to-section assignment and assignment history.",
    "APPLICATIONS": "Immutable applicant snapshot and current workflow state.",
    "APPLICATION_DOCUMENTS": "Uploaded application requirement metadata and review state.",
    "APPLICATION_STATUS_HISTORY": "Append-only application workflow history.",
    "STUDENTS": "Permanent student identity, independent of current placement.",
    "GUARDIANS": "Reusable guardian identity and contact record.",
    "STUDENT_GUARDIANS": "Student-to-guardian relationship, authority, and contact flags.",
    "STUDENT_DOCUMENTS": "Post-enrollment student document metadata and verification state.",
    "ENROLLMENTS": "A student's grade and section placement for one school year.",
    "ENROLLMENT_STATUS_HISTORY": "Append-only enrollment, transfer, and section-change history.",
    "RFID_CARDS": "Physical RFID credential inventory.",
    "RFID_ASSIGNMENTS": "Time-bounded RFID ownership history.",
    "RFID_DEVICES": "Registered RFID readers, gates, and device authentication state.",
    "ATTENDANCE_RECORDS": "Resolved daily attendance summary for an enrollment.",
    "RFID_SCAN_EVENTS": "Immutable raw RFID device events for audit and reprocessing.",
    "ATTENDANCE_CORRECTIONS": "Append-only manual attendance correction history.",
    "SUBJECTS": "Subject catalog for academic grading.",
    "GRADING_PERIODS": "Quarter, trimester, or semester definitions for a school year.",
    "CLASS_OFFERINGS": "A subject taught to a section by a teacher.",
    "STUDENT_GRADES": "Student mark for one class offering and grading period.",
    "AUDIT_LOGS": "Append-only security and business change trail.",
}

MODULES = [
    ("Identity and Access", ["USERS", "STAFF_PROFILES", "POSITIONS", "DEPARTMENTS", "AUDIT_LOGS"]),
    ("Academic Structure", ["SCHOOL_YEARS", "GRADE_LEVELS", "SECTIONS", "TEACHERS", "SECTION_TEACHERS"]),
    ("Applications and Enrollment", ["APPLICATIONS", "APPLICATION_DOCUMENTS", "APPLICATION_STATUS_HISTORY", "STUDENTS", "GUARDIANS", "STUDENT_GUARDIANS", "STUDENT_DOCUMENTS", "ENROLLMENTS", "ENROLLMENT_STATUS_HISTORY"]),
    ("RFID and Attendance", ["RFID_CARDS", "RFID_ASSIGNMENTS", "RFID_DEVICES", "ATTENDANCE_RECORDS", "RFID_SCAN_EVENTS", "ATTENDANCE_CORRECTIONS"]),
    ("Academic Grading", ["SUBJECTS", "GRADING_PERIODS", "CLASS_OFFERINGS", "STUDENT_GRADES"]),
]

OPTIONAL_NAMES = {
    "email_verified_at", "last_login_at", "password_changed_at", "remember_token", "deleted_at",
    "middle_name", "suffix", "birth_date", "gender", "mobile", "address", "department_id",
    "position_id", "photo_path", "hire_date", "user_id", "source_application_id", "lrn",
    "previous_school", "guardian_name", "guardian_relationship", "guardian_contact", "guardian_email",
    "strand", "program", "submitted_at", "verified_at", "approved_at", "reviewed_by", "reviewed_at",
    "review_notes", "from_status", "changed_by", "remarks", "alternate_mobile", "section_id",
    "application_id", "enrolled_on", "created_by", "from_section_id", "to_section_id", "reason",
    "issued_at", "notes", "student_id", "teacher_id", "staff_profile_id", "assigned_by", "ends_at",
    "end_reason", "location", "secret_hash", "last_seen_at", "rfid_card_id", "attendance_record_id",
    "error_message", "first_time_in", "last_time_out", "recorded_by", "previous_time_in", "new_time_in",
    "previous_time_out", "new_time_out", "description", "units", "numeric_grade", "letter_grade",
    "encoded_by", "approved_by", "approved_at", "entity_id", "old_values", "new_values", "ip_address",
    "user_agent", "request_id", "ended_on",
}

COLUMN_PURPOSES = {
    "id": "Internal primary key.", "public_id": "ULID exposed through APIs instead of the numeric key.",
    "username": "Unique login name.", "email": "Email address; normalized before uniqueness checks.",
    "password": "One-way password hash; never plaintext.", "role": "Authorization role used by policies and middleware.",
    "is_active": "Controls whether the record/account is operational.", "created_at": "Record creation timestamp.",
    "updated_at": "Last modification timestamp.", "deleted_at": "Laravel soft-deletion timestamp.",
    "staff_no": "Human-readable staff identifier.", "student_no": "Official internal student number.",
    "teacher_no": "Official teacher identifier.", "application_no": "Human-readable application reference.",
    "enrollment_no": "Human-readable enrollment reference.", "first_name": "Given name.",
    "middle_name": "Optional middle name.", "last_name": "Family name.", "suffix": "Optional name suffix.",
    "birth_date": "Date of birth; age is calculated, not stored.", "mobile": "Contact number stored as text.",
    "address": "Mailing or home address.", "code": "Stable machine/business code.", "name": "Human-readable name.",
    "description": "Optional explanatory text.", "status": "Current lifecycle or workflow state.",
    "school_year_id": "References the applicable school year.", "grade_level_id": "References the education level.",
    "section_id": "References the assigned class section.", "department_id": "References a controlled department.",
    "position_id": "References a controlled position.", "user_id": "Optional or required linked login account.",
    "student_id": "References a student.", "teacher_id": "References a teacher.",
    "guardian_id": "References a guardian.", "application_id": "References an enrollment application.",
    "enrollment_id": "References the student's school-year enrollment.", "rfid_card_id": "References a physical RFID card.",
    "rfid_device_id": "References the scanning device.", "attendance_record_id": "References the resolved daily attendance row.",
    "uid": "Normalized, unique RFID hardware identifier.", "event_uuid": "Idempotency key supplied by the device.",
    "attendance_date": "School-calendar date for the attendance summary.", "first_time_in": "Earliest accepted entry time.",
    "last_time_out": "Latest accepted exit time.", "source": "Origin of the record: RFID, manual, import, or system.",
    "changed_at": "Exact time of an append-only workflow transition.", "from_status": "Workflow state before the change.",
    "to_status": "Workflow state after the change.", "capacity": "Maximum permitted active enrollment count.",
    "sort_order": "Logical display and progression order.", "is_current": "Marks the operational school year.",
    "starts_on": "Start date of the represented period.", "ends_on": "End date of the represented period.",
    "storage_path": "Generated path in protected file storage.", "original_name": "Original upload filename for display only.",
    "mime_type": "Verified file media type.", "size_bytes": "Exact file size in bytes.",
    "sha256": "File integrity and duplicate-detection hash.", "review_status": "Document review outcome.",
    "verification_status": "Document verification outcome.", "employment_status": "Staff employment lifecycle state.",
    "lrn": "Philippine Learner Reference Number.", "relationship": "Relationship between student and guardian.",
    "numeric_grade": "Exact numeric mark stored with decimal precision.", "letter_grade": "Optional letter equivalent.",
    "entity_type": "Type of audited entity.", "entity_id": "Identifier of audited entity.",
    "old_values": "Selected values before the audited change.", "new_values": "Selected values after the audited change.",
}


def parse_schema():
    text = SOURCE.read_text(encoding="utf-8")
    mermaid = re.search(r"```mermaid\s+(.*?)```", text, re.S).group(1)
    tables = {}
    for name, body in re.findall(r"^\s{4}([A-Z_]+) \{\s*(.*?)^\s{4}\}", mermaid, re.S | re.M):
        columns = []
        for line in body.splitlines():
            parts = line.strip().split()
            if len(parts) < 2:
                continue
            dtype, col = parts[0], parts[1]
            key = " ".join(parts[2:]).replace(",", ", ") if len(parts) > 2 else ""
            columns.append((col, dtype, key))
        tables[name] = columns
    relationships = []
    for left, symbol, right, label in re.findall(r'^\s{4}([A-Z_]+) (\|\|--o\||\|\|--o\{|\|\|--\|\{|o\|--o\{) ([A-Z_]+) : "?([^"\n]+)"?$', mermaid, re.M):
        relationships.append((left, symbol, right, label.strip()))
    return tables, relationships


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def font_run(run, size=9, bold=False, color=TEXT, italic=False, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def format_table(table, header=True, font_size=8.4):
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0 and header:
                set_cell_shading(cell, PALE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1.5)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    font_run(run, font_size, bold=(r_idx == 0 and header), color=NAVY if r_idx == 0 else TEXT)
    if header:
        set_repeat_table_header(table.rows[0])


def add_table(doc, headers, rows, widths, font_size=8.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(data):
            cells[idx].text = str(value)
    set_table_geometry(table, widths)
    format_table(table, font_size=font_size)
    return table


def column_purpose(name):
    if name in COLUMN_PURPOSES:
        return COLUMN_PURPOSES[name]
    if name.endswith("_id"):
        return "Foreign key to the related " + name[:-3].replace("_", " ") + " record."
    if name.startswith("is_") or name.startswith("can_") or name.startswith("has_"):
        return "Boolean business rule or permission flag."
    if name.endswith("_at"):
        return "Timestamp for this lifecycle event."
    if name.endswith("_on") or name.endswith("_date"):
        return "Calendar date for this business event."
    return name.replace("_", " ").capitalize() + "."


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    font_run(run, 8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("SPRYtech School Management System  |  Database Reference")
    font_run(run, 8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc, table_count, relationship_count):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(95)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DATABASE SCHEMA")
    font_run(r, 28, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("School Management System")
    font_run(r, 17, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(50)
    r = p.add_run("Normalized database dictionary, relationships, constraints, and implementation cheat sheets")
    font_run(r, 11, italic=True, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    for idx, (value, label) in enumerate(((table_count, "tables"), (relationship_count, "relationships"), ("MySQL 8", "target"))):
        if idx:
            r = p.add_run("   |   ")
            font_run(r, 11, color=GRID)
        r = p.add_run(f"{value} {label}")
        font_run(r, 11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(60)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Prepared from the revamped project schema  |  August 2026")
    font_run(r, 9.5, color=MUTED)
    doc.add_page_break()


def build():
    tables, relationships = parse_schema()
    doc = Document()
    configure_document(doc)
    add_cover(doc, len(tables), len(relationships))

    doc.add_heading("How to use this reference", level=1)
    doc.add_paragraph("This document is the implementation reference for Laravel migrations, Eloquent relationships, API validation, and database review. Requiredness and delete behavior should be enforced at both the database and service layers.")
    add_table(doc, ["Marker", "Meaning"], [
        ("PK", "Primary key"), ("FK", "Foreign key"), ("UK", "Unique key"),
        ("NULL", "Optional value"), ("1:N", "One parent may have many children"),
        ("1:1", "At most one related row"), ("N:N", "Many-to-many through a junction table"),
    ], [1700, 7660], 9)

    doc.add_heading("Architecture at a glance", level=1)
    for module, names in MODULES:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        r = p.add_run(module + ": ")
        font_run(r, 10.5, bold=True, color=NAVY)
        r = p.add_run(", ".join(name.lower() for name in names))
        font_run(r, 10.5)

    doc.add_heading("Normalization principles", level=2)
    for item in (
        "users stores authentication only; identity and employment data belong in profile tables.",
        "Grade and section placement belongs in enrollments so academic history is retained.",
        "Applications are submission snapshots; workflow transitions are recorded in append-only history.",
        "RFID cards are independent assets, with time-bounded owner assignments.",
        "Raw RFID scans and resolved daily attendance are separate layers.",
        "Uploaded binaries remain in protected storage; the database stores metadata and hashes.",
    ):
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)

    doc.add_page_break()
    doc.add_heading("Complete table dictionary", level=1)
    doc.add_paragraph("Columns below reflect the complete proposed schema. FK columns are indexed. Composite uniqueness and cross-table business constraints are summarized after the dictionary.")

    table_number = 1
    for module, names in MODULES:
        doc.add_heading(module, level=1)
        for name in names:
            columns = tables.get(name, [])
            doc.add_heading(name.lower(), level=2)
            p = doc.add_paragraph()
            r = p.add_run("Purpose. ")
            font_run(r, 10, bold=True, color=NAVY)
            r = p.add_run(TABLE_PURPOSES.get(name, "Database entity."))
            font_run(r, 10)
            rows = []
            for col, dtype, key in columns:
                nullable = "Yes" if col in OPTIONAL_NAMES else "No"
                rows.append((col, dtype, key or "-", nullable, column_purpose(col)))
            add_table(doc, ["Column", "Data type", "Key", "Null", "Significance"], rows, [1850, 1550, 800, 620, 4540], 8.0)
            p = doc.add_paragraph(f"Table {table_number}. {name.lower()} column definition")
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(7)
            for run in p.runs:
                font_run(run, 8, italic=True, color=MUTED)
            table_number += 1

    doc.add_page_break()
    doc.add_heading("Relationship catalog", level=1)
    rel_rows = []
    symbol_map = {"||--o|": "1 : 0..1", "||--o{": "1 : N", "||--|{": "1 : 1..N", "o|--o{": "0..1 : N"}
    for left, symbol, right, label in relationships:
        rel_rows.append((left.lower(), symbol_map.get(symbol, symbol), right.lower(), label))
    add_table(doc, ["Parent", "Cardinality", "Child", "Relationship"], rel_rows, [2300, 1350, 2300, 3410], 8.2)

    doc.add_heading("Foreign-key constraints cheat sheet", level=1)
    fk_rows = []
    for table_name, columns in tables.items():
        for col, _, key in columns:
            if "FK" not in key:
                continue
            target = col[:-3] if col.endswith("_id") else col
            target_table = {
                "user": "users", "department": "departments", "position": "positions", "school_year": "school_years",
                "grade_level": "grade_levels", "section": "sections", "teacher": "teachers", "application": "applications",
                "reviewed_by": "users", "changed_by": "users", "student": "students", "guardian": "guardians",
                "verified_by": "users", "created_by": "users", "enrollment": "enrollments", "from_section": "sections",
                "to_section": "sections", "rfid_card": "rfid_cards", "staff_profile": "staff_profiles", "assigned_by": "users",
                "rfid_device": "rfid_devices", "attendance_record": "attendance_records", "recorded_by": "users",
                "corrected_by": "users", "subject": "subjects", "class_offering": "class_offerings",
                "grading_period": "grading_periods", "encoded_by": "users", "approved_by": "users",
            }.get(target, target + "s")
            optional = col in OPTIONAL_NAMES
            if col.endswith("_by") or col in {"user_id", "department_id", "position_id", "section_id", "application_id", "rfid_card_id", "attendance_record_id", "student_id", "teacher_id", "staff_profile_id", "from_section_id", "to_section_id"}:
                on_delete = "SET NULL" if optional else "RESTRICT"
            elif table_name in {"APPLICATION_DOCUMENTS", "APPLICATION_STATUS_HISTORY", "STUDENT_GUARDIANS", "ENROLLMENT_STATUS_HISTORY"}:
                on_delete = "CASCADE"
            else:
                on_delete = "RESTRICT"
            fk_rows.append((table_name.lower(), col, f"{target_table}.id", on_delete, "Optional" if optional else "Required"))
    add_table(doc, ["Child table", "FK column", "References", "On delete", "Requiredness"], fk_rows, [2050, 1900, 2050, 1450, 1910], 7.9)

    doc.add_heading("Composite uniqueness and checks", level=1)
    constraints = [
        ("grade_levels", "UNIQUE(department_id, name)", "Prevents duplicate grade names within a department."),
        ("sections", "UNIQUE(school_year_id, grade_level_id, name)", "Prevents duplicate sections for the same grade and year."),
        ("sections", "CHECK(capacity > 0)", "Rejects invalid class capacity."),
        ("student_guardians", "UNIQUE(student_id, guardian_id)", "Prevents repeated guardian links."),
        ("application_documents", "UNIQUE(application_id, document_type)", "One current upload per requirement type."),
        ("enrollments", "UNIQUE(student_id, school_year_id)", "One enrollment per student per school year."),
        ("attendance_records", "UNIQUE(enrollment_id, attendance_date)", "One daily attendance summary."),
        ("rfid_assignments", "CHECK(exactly one owner FK is populated)", "Prevents ambiguous card ownership."),
        ("rfid_scan_events", "UNIQUE(event_uuid)", "Makes device submissions idempotent."),
        ("class_offerings", "UNIQUE(section_id, subject_id)", "Prevents duplicate subject offerings."),
        ("student_grades", "UNIQUE(enrollment_id, class_offering_id, grading_period_id)", "One mark per period and offering."),
    ]
    add_table(doc, ["Table", "Constraint", "Reason"], constraints, [1900, 3400, 4060], 8.2)

    doc.add_heading("Service-layer integrity rules", level=1)
    for item in (
        "Lock the section row while checking capacity and creating an enrollment.",
        "Ensure an enrollment's section belongs to the same grade level and school year.",
        "Allow only one current school year.",
        "Allow only one active RFID assignment per card and per owner.",
        "Reject scans from inactive devices, inactive cards, or cards without an active assignment.",
        "Prevent changes to closed school years and locked grading periods without an authorized reopening workflow.",
        "Create student, enrollment, status history, and audit records in one transaction when enrolling an approved applicant.",
    ):
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)

    doc.add_heading("Laravel migration cheat sheet", level=1)
    migration_rows = [
        ("Primary key", "$table->id();", "BIGINT UNSIGNED AUTO_INCREMENT"),
        ("Foreign key", "$table->foreignId('student_id')->constrained();", "Index + FK to students.id"),
        ("Nullable FK", "$table->foreignId('section_id')->nullable()->constrained();", "Optional relationship"),
        ("Restrict delete", "->restrictOnDelete();", "Preserve referenced history"),
        ("Set null", "->nullOnDelete();", "Retain row after optional actor/parent removal"),
        ("Cascade delete", "->cascadeOnDelete();", "Delete true dependent children"),
        ("Composite unique", "$table->unique(['student_id', 'school_year_id']);", "Business-key enforcement"),
        ("Soft delete", "$table->softDeletes();", "Adds deleted_at"),
        ("Exact decimal", "$table->decimal('numeric_grade', 5, 2);", "Avoid floating-point marks"),
        ("File size", "$table->unsignedBigInteger('size_bytes');", "Exact byte count"),
    ]
    add_table(doc, ["Pattern", "Laravel example", "Meaning"], migration_rows, [1700, 4600, 3060], 8.1)

    doc.add_heading("Indexing cheat sheet", level=1)
    index_rows = [
        ("users", "(role, is_active)", "Role-filtered account lists"),
        ("staff_profiles", "(last_name, first_name)", "Staff search"),
        ("teachers", "(last_name, first_name)", "Teacher search"),
        ("students", "(last_name, first_name)", "Student search"),
        ("applications", "(status, submitted_at)", "Verification queue"),
        ("enrollments", "(section_id, status)", "Class rosters and capacity"),
        ("attendance_records", "(attendance_date, status)", "Daily attendance reporting"),
        ("rfid_assignments", "(rfid_card_id, ends_at)", "Resolve current owner"),
        ("rfid_scan_events", "(rfid_card_id, scanned_at)", "Card scan history"),
        ("audit_logs", "(entity_type, entity_id)", "Entity audit trail"),
    ]
    add_table(doc, ["Table", "Index", "Supports"], index_rows, [2200, 3100, 4060], 8.3)

    doc.add_heading("Status vocabulary", level=1)
    status_rows = [
        ("User account", "active / inactive", "users.is_active"),
        ("Employment", "active / inactive / on_leave / separated", "staff_profiles.employment_status"),
        ("Application", "draft / submitted / under_review / verified / approved / waitlisted / rejected / enrolled / cancelled", "applications.status"),
        ("Enrollment", "pending / enrolled / completed / withdrawn / transferred / cancelled", "enrollments.status"),
        ("RFID card", "active / inactive / lost / damaged / retired", "rfid_cards.status"),
        ("Attendance", "present / absent / late / excused / half_day", "attendance_records.status"),
        ("Grade", "draft / submitted / approved / locked", "student_grades.status"),
    ]
    add_table(doc, ["Domain", "Allowed values", "Stored in"], status_rows, [1700, 5000, 2660], 8.1)

    doc.add_heading("Final implementation notes", level=1)
    doc.add_paragraph("Use lowercase machine values in the database, UTC timestamps, DATE for calendar dates, VARCHAR for phone numbers and RFID UIDs, DECIMAL for marks, and private storage for uploaded documents. Frontend validation improves usability; Laravel Form Requests and database constraints remain the authoritative protection layers.")

    props = doc.core_properties
    props.title = "School Management System Database Schema"
    props.subject = "Normalized schema, relationships, foreign-key constraints, and implementation cheat sheets"
    props.author = "SPRYtech Development Team"
    props.keywords = "database schema, Laravel, MySQL, school management system"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
